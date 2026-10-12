"""将业务角色范围表同步为项目的 SID-管理员角色关键词配置表。"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_SID_COLUMN = "SID"
SOURCE_ROLE_COLUMN = "根据角色-权限矩阵或系统查看确认的应用层管理员角色范围"
TARGET_HEADERS = ["SID", "管理员角色关键词"]
NO_ROLE_VALUES = {"", "N/A", "NA", "TBD"}


def split_role_keywords(value: object) -> list[str]:
    """按业务表中的换行及中英文分隔符拆成一个角色一个关键词。"""
    text = "" if value is None else str(value).strip()
    if text.upper() in NO_ROLE_VALUES:
        return []
    return [part.strip() for part in re.split(r"[\n,，、;；]+", text) if part.strip()]


def build_rows(source_path: Path) -> tuple[list[list[str]], list[str]]:
    df = pd.read_excel(source_path, dtype=str, keep_default_na=False)
    missing = {SOURCE_SID_COLUMN, SOURCE_ROLE_COLUMN} - set(df.columns)
    if missing:
        raise ValueError(f"{source_path} 缺少字段：{sorted(missing)}")

    rows: list[list[str]] = []
    skipped_sids: list[str] = []
    seen: set[tuple[str, str]] = set()
    for _, record in df.iterrows():
        sid = str(record[SOURCE_SID_COLUMN]).strip()
        keywords = split_role_keywords(record[SOURCE_ROLE_COLUMN])
        if not sid:
            continue
        if not keywords:
            skipped_sids.append(sid)
            continue
        for keyword in keywords:
            key = (sid.upper(), keyword.upper())
            if key not in seen:
                rows.append([sid, keyword])
                seen.add(key)
    return rows, skipped_sids


def write_config(rows: list[list[str]], target_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "应用层管理员角色配置"
    ws.append(TARGET_HEADERS)
    for row in rows:
        ws.append(row)

    header = ws[1]
    for cell in header:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for column in ("A", "B"):
        for cell in ws[column]:
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 42
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{max(ws.max_row, 1)}"
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="同步 SID 应用层管理员角色关键词配置")
    parser.add_argument("--source", required=True, help="各系统应用层管理员角色定义范围.xlsx 路径")
    parser.add_argument(
        "--target",
        default="configs/应用层管理员角色配置表.xlsx",
        help="项目配置表输出路径",
    )
    args = parser.parse_args()

    rows, skipped_sids = build_rows(Path(args.source))
    write_config(rows, Path(args.target))
    print(f"已写入 {len(rows)} 条 SID-管理员角色关键词：{args.target}")
    if skipped_sids:
        print("未写入角色关键词（来源为空/N/A/TBD）：" + " | ".join(skipped_sids))


if __name__ == "__main__":
    main()
